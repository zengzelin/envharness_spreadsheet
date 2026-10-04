from __future__ import annotations

from pathlib import Path
import re
import zipfile

import openpyxl
from openpyxl.styles import Font

from envharness.bridges.spreadsheetbench.badcase_diagnostics import (
    DiagnosticLimits,
    diagnose_failed_workbook,
)


def _write_workbook(path: Path, sheets: dict[str, dict[str, object]]) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for sheet_name, cells in sheets.items():
        sheet = workbook.create_sheet(sheet_name)
        for coordinate, value in cells.items():
            sheet[coordinate] = value
    workbook.save(path)
    workbook.close()


def _set_cached_formula_value(path: Path, coordinate: str, value: object) -> None:
    """Inject a cached result so data_only=True behaves like a recalc'd file."""
    temporary = path.with_suffix(".cached.xlsx")
    pattern = re.compile(
        rf'(<c\s+r="{re.escape(coordinate)}"[^>]*>.*?<f[^>]*>.*?</f>\s*<v>).*?(</v>)'
    )
    replaced = False
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary, "w", zipfile.ZIP_DEFLATED
    ) as destination:
        for member in source.infolist():
            payload = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                xml = payload.decode("utf-8")
                xml, count = pattern.subn(
                    lambda match: match.group(1) + str(value) + match.group(2),
                    xml,
                    count=1,
                )
                replaced = replaced or count == 1
                payload = xml.encode("utf-8")
            destination.writestr(member, payload)
    assert replaced, f"formula cell {coordinate} was not found in worksheet XML"
    temporary.replace(path)


def test_diagnostics_counts_answer_matches_and_bounds_examples(
    tmp_path: Path,
) -> None:
    input_path = tmp_path / "input.xlsx"
    golden_path = tmp_path / "golden.xlsx"
    output_path = tmp_path / "output.xlsx"
    _write_workbook(input_path, {"Sheet1": {"A1": 0, "A2": 0, "A3": 0}})
    _write_workbook(golden_path, {"Sheet1": {"A1": 1, "A2": 2, "A3": 3}})
    _write_workbook(
        output_path,
        {"Sheet1": {"A1": 1, "A2": 9, "A3": "x" * 400}},
    )

    result = diagnose_failed_workbook(
        str(input_path),
        str(golden_path),
        str(output_path),
        "Sheet1!A1:A3",
        limits=DiagnosticLimits(max_scan_cells=100, max_examples=1),
    )

    assert result["status"] == "complete"
    assert result["answer_cells_total"] == 3
    assert result["answer_cells_scanned"] == 3
    assert result["answer_cells_matched"] == 1
    assert result["answer_cells_mismatched"] == 2
    assert result["answer_match_ratio"] == 1 / 3
    assert result["answer_cells_changed_from_input"] == 3
    assert len(result["first_mismatches"]) == 1
    assert all(
        len(value) <= 200
        for value in result["first_mismatches"][0].values()
        if isinstance(value, str)
    )


def test_diagnostics_detects_untouched_answer_and_mutations_on_other_sheets(
    tmp_path: Path,
) -> None:
    input_path = tmp_path / "input.xlsx"
    golden_path = tmp_path / "golden.xlsx"
    output_path = tmp_path / "output.xlsx"
    _write_workbook(
        input_path,
        {"Sheet1": {"A1": 0, "B1": 0}, "Other": {"A1": 0}},
    )
    _write_workbook(
        golden_path,
        {"Sheet1": {"A1": 1, "B1": 0}, "Other": {"A1": 0}},
    )
    _write_workbook(
        output_path,
        {"Sheet1": {"A1": 0, "B1": 2}, "Other": {"A1": 3}},
    )

    result = diagnose_failed_workbook(
        str(input_path),
        str(golden_path),
        str(output_path),
        "Sheet1!A1",
        limits=DiagnosticLimits(max_scan_cells=100, max_examples=20),
    )

    assert result["answer_range_untouched"] is True
    assert result["answer_required_change_count"] == 1
    assert result["outside_answer_cells_changed"] == 2
    assert result["mutated_sheets"] == ["Other", "Sheet1"]
    assert result["diagnostic_truncated"] is False


def test_diagnostics_uses_one_shared_coordinate_budget(tmp_path: Path) -> None:
    input_path = tmp_path / "input.xlsx"
    golden_path = tmp_path / "golden.xlsx"
    output_path = tmp_path / "output.xlsx"
    cells = {"A1": 0, "A2": 0, "A3": 0, "A4": 0}
    _write_workbook(input_path, {"Sheet1": cells})
    _write_workbook(golden_path, {"Sheet1": {**cells, "A1": 1}})
    _write_workbook(output_path, {"Sheet1": {**cells, "A4": 4}})

    result = diagnose_failed_workbook(
        str(input_path),
        str(golden_path),
        str(output_path),
        "Sheet1!A1:A2",
        limits=DiagnosticLimits(max_scan_cells=3, max_examples=20),
    )

    assert result["answer_cells_scanned"] == 2
    assert result["diagnostic_cells_scanned"] == 3
    assert result["diagnostic_truncated"] is True


def test_diagnostics_distinguishes_formula_structure_results_and_style(
    tmp_path: Path,
) -> None:
    input_path = tmp_path / "input.xlsx"
    golden_path = tmp_path / "golden.xlsx"
    output_path = tmp_path / "output.xlsx"
    _write_workbook(
        input_path,
        {"Sheet1": {"A1": 0, "A2": 0, "A3": 0, "A4": 1}},
    )
    _write_workbook(
        golden_path,
        {"Sheet1": {"A1": "=1+1", "A2": "=1+1", "A3": "=1+1", "A4": 1}},
    )
    golden = openpyxl.load_workbook(golden_path)
    golden["Sheet1"]["A4"].font = Font(bold=True)
    golden.save(golden_path)
    golden.close()
    _set_cached_formula_value(golden_path, "A1", 2)
    _set_cached_formula_value(golden_path, "A2", 2)
    _set_cached_formula_value(golden_path, "A3", 2)
    _write_workbook(
        output_path,
        {"Sheet1": {"A1": 2, "A2": "=1+2", "A3": "=2", "A4": 1}},
    )
    _set_cached_formula_value(output_path, "A2", 3)
    _set_cached_formula_value(output_path, "A3", 2)

    result = diagnose_failed_workbook(
        str(input_path),
        str(golden_path),
        str(output_path),
        "Sheet1!A1:A4",
        limits=DiagnosticLimits(max_scan_cells=100, max_examples=20),
    )

    assert result["formula_cells_expected"] == 3
    assert result["formula_cells_output"] == 2
    assert result["formula_to_static_count"] == 1
    assert result["formula_result_mismatch_count"] == 1
    assert result["format_mismatch_count"] == 1


def test_diagnostics_reports_missing_and_unreadable_output(tmp_path: Path) -> None:
    input_path = tmp_path / "input.xlsx"
    golden_path = tmp_path / "golden.xlsx"
    output_path = tmp_path / "output.xlsx"
    _write_workbook(input_path, {"Sheet1": {"A1": 0}})
    _write_workbook(golden_path, {"Sheet1": {"A1": 1}})

    missing = diagnose_failed_workbook(
        str(input_path),
        str(golden_path),
        str(output_path),
        "Sheet1!A1",
        limits=DiagnosticLimits(max_scan_cells=100, max_examples=20),
    )
    assert missing["status"] == "output_missing"

    output_path.write_text("not an xlsx", encoding="utf-8")
    unreadable = diagnose_failed_workbook(
        str(input_path),
        str(golden_path),
        str(output_path),
        "Sheet1!A1",
        limits=DiagnosticLimits(max_scan_cells=100, max_examples=20),
    )
    assert unreadable["status"] == "load_error"
    assert unreadable["diagnostic_error_type"]
