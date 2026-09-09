from __future__ import annotations

from pathlib import Path

import openpyxl

from envharness.bridges.spreadsheetbench import online_judge_eval


def test_compare_workbooks_accepts_verified_400_answer_position_variants(
    tmp_path: Path,
) -> None:
    golden = tmp_path / "golden.xlsx"
    processed = tmp_path / "processed.xlsx"
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    sheet3 = workbook.create_sheet("Sheet3")
    sheet4 = workbook.create_sheet("Sheet4")
    sheet1 = workbook.create_sheet("Sheet1")
    comma_sheet = workbook.create_sheet("b2b, sez, de")
    for sheet in (sheet3, sheet4):
        sheet.append(["a", "b", "c", "d", "e", "f", "g"])
        sheet.append([1, 2, 3, 4, 5, 6, 7])
    for _ in range(4):
        comma_sheet.append([None] * 22)
    comma_sheet.append(list(range(1, 23)))
    comma_sheet.append(list(range(23, 45)))
    for row in range(2, 7):
        sheet1.cell(row=row, column=56, value=row)
    workbook.save(golden)
    workbook.save(processed)

    passed, msg = online_judge_eval.compare_workbooks(
        str(golden),
        str(processed),
        "Sheet-Level Manipulation",
        "Sheet3'!A:G,'Sheet4'!A:G,'b2b, sez, de'!A5:V6,'Sheet1'!BD2:6",
    )

    assert passed is True
    assert msg == ""
