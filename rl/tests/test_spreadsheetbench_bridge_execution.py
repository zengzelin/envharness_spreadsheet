from __future__ import annotations

import json
from pathlib import Path
import time

import openpyxl

from envharness.bridges.spreadsheetbench import write_tools
from envharness.bridges.spreadsheetbench.bridge import (
    SpreadsheetBenchEnv,
    SpreadsheetBenchEnvState,
)
from envharness.core.types import Action


def _execution_env(tmp_path: Path) -> SpreadsheetBenchEnv:
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "output.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Sheet1"
    workbook["Sheet1"]["A1"] = "Name"
    workbook["Sheet1"]["A2"] = "Alpha"
    workbook["Sheet1"]["B2"] = "=1+1"
    second = workbook.create_sheet("Hidden Data")
    second.sheet_state = "hidden"
    second["C3"] = "Needle"
    workbook.save(input_path)
    workbook.save(output_path)

    env = SpreadsheetBenchEnv()
    env._tool_set = "native_read"
    env._workdir = str(tmp_path)
    env.state = SpreadsheetBenchEnvState(
        task_id="task-1",
        instruction="Write A1",
        answer_position="'Sheet1'!A1",
        answer_sheet="Sheet1",
        input_path=str(input_path),
        output_path=str(output_path),
    )
    return env


def _tool_payload(response) -> dict:
    marker = "last tool output:\n"
    assert marker in response.observation.text
    return json.loads(response.observation.text.split(marker, 1)[1])


def test_run_python_predefines_paths_and_workbook_helpers(tmp_path: Path) -> None:
    env = _execution_env(tmp_path)

    response = env.step(Action(name="run_python", kwargs={
        "code": (
            "wb = load_workbook_for_edit()\n"
            "wb['Sheet1']['A1'] = input_path\n"
            "save_workbook(wb)"
        )
    }))

    workbook = openpyxl.load_workbook(env.state.output_path)
    assert workbook["Sheet1"]["A1"].value == env.state.input_path
    assert response.reward == 0.0
    assert response.info["returncode"] == 0
    assert response.info["python_error"] is False
    assert response.info["python_error_type"] == ""


def test_run_python_returns_structured_syntax_error_and_penalty(
    tmp_path: Path,
) -> None:
    env = _execution_env(tmp_path)
    env._python_error_penalty = -0.05
    env._syntax_error_penalty = -0.1

    response = env.step(Action(name="run_python", kwargs={
        "code": "for value in:\n    pass"
    }))

    assert response.reward == -0.1
    assert response.info["returncode"] != 0
    assert response.info["python_error"] is True
    assert response.info["python_error_type"] == "SyntaxError"
    assert response.info["syntax_error"] is True


def test_run_python_classifies_exception_before_output_truncation(
    tmp_path: Path,
) -> None:
    env = _execution_env(tmp_path)
    env._obs_truncate = 100

    response = env.step(Action(name="run_python", kwargs={
        "code": "print('x' * 500)\nraise ValueError('bad data')"
    }))

    assert response.info["python_error_type"] == "ValueError"
    assert response.observation.text.endswith("...[output truncated]")


def test_run_python_timeout_reaps_child_process_output_pipes(tmp_path: Path) -> None:
    env = _execution_env(tmp_path)
    env._step_timeout = 0.2
    started = time.monotonic()

    response = env.step(Action(name="run_python", kwargs={
        "code": (
            "import subprocess, sys, time\n"
            "subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(2)'], "
            "stdout=sys.stdout, stderr=sys.stderr)\n"
            "time.sleep(5)\n"
        )
    }))

    assert time.monotonic() - started < 1.5
    assert response.info["python_error_type"] == "TimeoutExpired"


def test_run_python_logs_task_and_stage(capsys, tmp_path: Path) -> None:
    env = _execution_env(tmp_path)

    env.step(Action(name="run_python", kwargs={"code": "print('ok')"}))

    output = capsys.readouterr().out
    assert "task_id=task-1 step=1 action=run_python stage=run_python START" in output
    assert "task_id=task-1 step=1 action=run_python stage=run_python END" in output


def test_validate_workbook_checks_output_without_grading(tmp_path: Path) -> None:
    env = _execution_env(tmp_path)

    response = env.step(Action(name="validate_workbook", kwargs={}))

    assert response.reward == 0.0
    assert response.terminated is False
    assert response.info["validation_ok"] is True
    assert "Sheet1" in response.observation.text


def test_validate_workbook_scans_large_range_sequentially(
    tmp_path: Path, monkeypatch,
) -> None:
    env = _execution_env(tmp_path)
    workbook = openpyxl.load_workbook(env.state.output_path)
    sheet = workbook["Sheet1"]
    sheet["J10411"] = "last"
    workbook.save(env.state.output_path)
    env.state.answer_position = "'Sheet1'!A1:J10411"

    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    def fail_random_access(self, key):
        raise AssertionError(f"read-only random cell access: {key}")

    monkeypatch.setattr(ReadOnlyWorksheet, "__getitem__", fail_random_access)
    response = env.step(Action(name="validate_workbook", kwargs={}))

    assert response.info["validation_ok"] is True
    assert "cells=104110, nonempty=4, formulas=1" in response.observation.text


def test_list_sheets_reports_order_visibility_and_dimensions(tmp_path: Path) -> None:
    env = _execution_env(tmp_path)

    response = env.step(Action(name="list_sheets", kwargs={}))

    payload = _tool_payload(response)
    assert response.reward == 0.0
    assert response.info["tool_ok"] is True
    assert payload["status"] == "success"
    assert payload["sheet_count"] == 2
    assert payload["sheets"][0] == {
        "name": "Sheet1",
        "index": 0,
        "state": "visible",
        "active": True,
        "max_row": 2,
        "max_column": 2,
    }
    assert payload["sheets"][1]["state"] == "hidden"


def test_inspect_range_returns_formulas_and_summary(tmp_path: Path) -> None:
    env = _execution_env(tmp_path)

    cells_response = env.step(Action(name="inspect_range", kwargs={
        "range": "A1:B2",
        "sheet_name": "Sheet1",
    }))
    summary_response = env.step(Action(name="inspect_range", kwargs={
        "range": "'Hidden Data'!A1:C3",
        "mode": "summary",
    }))

    cells = _tool_payload(cells_response)
    summary = _tool_payload(summary_response)
    assert cells["cells"][3]["address"] == "B2"
    assert cells["cells"][3]["formula"] == "=1+1"
    assert summary["summary"]["cells"] == 9
    assert summary["summary"]["non_empty_cells"] == 1
    assert summary["summary"]["formula_cells"] == 0


def test_find_cells_searches_values_and_formulas(tmp_path: Path) -> None:
    env = _execution_env(tmp_path)

    value_response = env.step(Action(name="find_cells", kwargs={
        "query": "alpha",
        "match": "equals",
        "case_sensitive": False,
        "include_values": True,
    }))
    formula_response = env.step(Action(name="find_cells", kwargs={
        "query": "1+1",
        "search_in": "formulas",
        "return_mode": "all",
    }))

    values = _tool_payload(value_response)
    formulas = _tool_payload(formula_response)
    assert values["matches"] == [{
        "sheet": "Sheet1", "address": "A2", "value": "Alpha"
    }]
    assert formulas["matches"] == [{
        "sheet": "Sheet1", "address": "B2", "formula": "=1+1"
    }]


def test_native_read_tools_are_disabled_in_python_tool_set(tmp_path: Path) -> None:
    env = _execution_env(tmp_path)
    env._tool_set = "python"

    response = env.step(Action(name="list_sheets", kwargs={}))

    assert response.info["error"] == "tool_disabled"
    assert "SPREADSHEETBENCH_TOOL_SET=native_read" in response.observation.text


def test_tool_schemas_follow_selected_tool_set(monkeypatch) -> None:
    monkeypatch.setenv("SPREADSHEETBENCH_TOOL_SET", "python")
    python_names = {
        item["function"]["name"] for item in SpreadsheetBenchEnv.tool_schemas()
    }
    monkeypatch.setenv("SPREADSHEETBENCH_TOOL_SET", "native_read")
    native_names = {
        item["function"]["name"] for item in SpreadsheetBenchEnv.tool_schemas()
    }
    monkeypatch.setenv("SPREADSHEETBENCH_TOOL_SET", "native_basic")
    basic_names = {
        item["function"]["name"] for item in SpreadsheetBenchEnv.tool_schemas()
    }

    assert python_names == {"run_python", "validate_workbook", "submit"}
    assert native_names == python_names | {
        "list_sheets", "inspect_range", "find_cells"
    }
    assert basic_names == native_names | {
        "write_range", "clear_range", "fill_formula"
    }
    write_schema = next(
        item for item in SpreadsheetBenchEnv.tool_schemas()
        if item["function"]["name"] == "write_range"
    )
    assert write_schema["function"]["parameters"]["required"] == ["range", "data"]
    assert write_schema["function"]["parameters"]["properties"]["data"]["type"] == [
        "array", "string", "number", "boolean"
    ]
    fill_schema = next(
        item for item in SpreadsheetBenchEnv.tool_schemas()
        if item["function"]["name"] == "fill_formula"
    )
    assert fill_schema["function"]["parameters"]["required"] == [
        "start_cell", "formula_template"
    ]


def test_native_read_tool_parameter_error_is_nonfatal(tmp_path: Path) -> None:
    env = _execution_env(tmp_path)

    response = env.step(Action(name="inspect_range", kwargs={
        "range": "A:A",
    }))

    payload = _tool_payload(response)
    assert response.terminated is False
    assert response.reward == 0.0
    assert response.info["tool_ok"] is False
    assert response.info["tool_error"] == "invalid_range"
    assert payload["status"] == "error"


def test_native_basic_write_range_updates_current_workbook(tmp_path: Path) -> None:
    env = _execution_env(tmp_path)
    env._tool_set = "native_basic"

    response = env.step(Action(name="write_range", kwargs={
        "range": "C2:D3",
        "sheet_name": "Sheet1",
        "data": [[1, 2], [3, 4]],
    }))

    workbook = openpyxl.load_workbook(env.state.output_path, data_only=False)
    assert response.info["tool_ok"] is True
    assert response.info["tool_category"] == "write"
    assert workbook["Sheet1"]["C2"].value == 1
    assert workbook["Sheet1"]["D3"].value == 4


def test_native_basic_write_range_accepts_column_data(tmp_path: Path) -> None:
    env = _execution_env(tmp_path)
    env._tool_set = "native_basic"

    response = env.step(Action(name="write_range", kwargs={
        "range": "C2:C4",
        "sheet_name": "Sheet1",
        "data": [1, 2, 3],
    }))

    workbook = openpyxl.load_workbook(env.state.output_path, data_only=False)
    assert response.info["tool_ok"] is True
    assert [workbook["Sheet1"][f"C{row}"].value for row in range(2, 5)] == [
        1, 2, 3
    ]


def test_native_basic_fill_formula_translates_relative_and_absolute_refs(
    tmp_path: Path,
) -> None:
    env = _execution_env(tmp_path)
    env._tool_set = "native_basic"

    response = env.step(Action(name="fill_formula", kwargs={
        "sheet_name": "Sheet1",
        "start_cell": "C2",
        "end_row": 4,
        "end_col": "D",
        "formula_template": "=A2+$B$1+C$1+$A2",
    }))

    workbook = openpyxl.load_workbook(env.state.output_path, data_only=False)
    assert response.info["tool_ok"] is True
    assert response.info["tool_category"] == "write"
    assert workbook["Sheet1"]["C2"].value == "=A2+$B$1+C$1+$A2"
    assert workbook["Sheet1"]["D2"].value == "=B2+$B$1+D$1+$A2"
    assert workbook["Sheet1"]["C4"].value == "=A4+$B$1+C$1+$A4"


def test_native_basic_fill_formula_rejects_non_formula_without_mutation(
    tmp_path: Path,
) -> None:
    env = _execution_env(tmp_path)
    env._tool_set = "native_basic"
    before = Path(env.state.output_path).read_bytes()

    response = env.step(Action(name="fill_formula", kwargs={
        "start_cell": "C2",
        "end_row": 4,
        "formula_template": "SUM(A1:A2)",
    }))

    assert response.info["tool_ok"] is False
    assert response.info["tool_error"] == "invalid_formula"
    assert Path(env.state.output_path).read_bytes() == before


def test_native_basic_fill_formula_rejects_malformed_formula_without_mutation(
    tmp_path: Path,
) -> None:
    env = _execution_env(tmp_path)
    env._tool_set = "native_basic"
    before = Path(env.state.output_path).read_bytes()

    response = env.step(Action(name="fill_formula", kwargs={
        "start_cell": "C2",
        "end_row": 4,
        "formula_template": "=SUM(A2:B2",
    }))

    assert response.info["tool_ok"] is False
    assert response.info["tool_error"] == "invalid_formula"
    assert "unclosed '('" in response.observation.text
    assert Path(env.state.output_path).read_bytes() == before


def test_native_basic_fill_formula_save_failure_preserves_workbook(
    tmp_path: Path, monkeypatch,
) -> None:
    env = _execution_env(tmp_path)
    env._tool_set = "native_basic"
    before = Path(env.state.output_path).read_bytes()

    def fail_save(workbook, path):
        raise RuntimeError("simulated save failure")

    monkeypatch.setattr(write_tools, "_atomic_save", fail_save)
    response = env.step(Action(name="fill_formula", kwargs={
        "start_cell": "C2",
        "end_row": 4,
        "formula_template": "=A2+B2",
    }))

    assert response.info["tool_ok"] is False
    assert response.info["tool_error"] == "tool_execution_failed"
    assert Path(env.state.output_path).read_bytes() == before


def test_native_basic_fill_formula_rejects_oversized_range(
    tmp_path: Path,
) -> None:
    env = _execution_env(tmp_path)
    env._tool_set = "native_basic"

    response = env.step(Action(name="fill_formula", kwargs={
        "start_cell": "A1",
        "end_row": 300001,
        "formula_template": "=B1",
    }))

    assert response.info["tool_ok"] is False
    assert response.info["tool_error"] == "range_too_large"


def test_native_basic_write_range_rejects_formulas_without_mutation(
    tmp_path: Path,
) -> None:
    env = _execution_env(tmp_path)
    env._tool_set = "native_basic"
    before = Path(env.state.output_path).read_bytes()

    response = env.step(Action(name="write_range", kwargs={
        "range": "A1",
        "data": "  =SUM(A2:A3)",
    }))

    assert response.info["tool_ok"] is False
    assert response.info["tool_error"] == "formula_not_allowed"
    assert Path(env.state.output_path).read_bytes() == before


def test_native_basic_write_range_rejects_null_top_level_data(
    tmp_path: Path,
) -> None:
    env = _execution_env(tmp_path)
    env._tool_set = "native_basic"

    response = env.step(Action(name="write_range", kwargs={
        "range": "A1",
        "data": None,
    }))

    assert response.info["tool_ok"] is False
    assert response.info["tool_error"] == "invalid_data"


def test_native_basic_write_range_runtime_failure_is_nonfatal(
    tmp_path: Path,
) -> None:
    env = _execution_env(tmp_path)
    env._tool_set = "native_basic"
    workbook = openpyxl.load_workbook(env.state.output_path)
    workbook["Sheet1"].merge_cells("C2:D2")
    workbook.save(env.state.output_path)

    response = env.step(Action(name="write_range", kwargs={
        "range": "D2",
        "data": "blocked",
    }))

    assert response.terminated is False
    assert response.info["tool_ok"] is False
    assert response.info["tool_error"] == "tool_execution_failed"


def test_native_basic_clear_range_clears_values_atomically(tmp_path: Path) -> None:
    env = _execution_env(tmp_path)
    env._tool_set = "native_basic"

    response = env.step(Action(name="clear_range", kwargs={
        "range": "A2:B2",
        "sheet_name": "Sheet1",
    }))

    workbook = openpyxl.load_workbook(env.state.output_path, data_only=False)
    assert response.info["tool_ok"] is True
    assert response.info["tool_category"] == "write"
    assert workbook["Sheet1"]["A2"].value is None
    assert workbook["Sheet1"]["B2"].value is None


def test_native_write_tools_are_disabled_in_native_read_mode(
    tmp_path: Path,
) -> None:
    env = _execution_env(tmp_path)

    response = env.step(Action(name="write_range", kwargs={
        "range": "A1",
        "data": "changed",
    }))

    assert response.info["error"] == "tool_disabled"
    workbook = openpyxl.load_workbook(env.state.output_path)
    assert workbook["Sheet1"]["A1"].value == "Name"
