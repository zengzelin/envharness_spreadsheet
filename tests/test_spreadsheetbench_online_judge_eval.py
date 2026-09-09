from __future__ import annotations

from pathlib import Path

import openpyxl

from envharness.bridges.spreadsheetbench import online_judge_eval


def test_generate_cell_names_expands_whole_column_range_against_sheet_bounds():
    cells = online_judge_eval._generate_cell_names("A:C", max_row=2)

    assert cells == ["A1", "A2", "B1", "B2", "C1", "C2"]


def test_generate_cell_names_infers_end_column_for_row_only_end_cell():
    cells = online_judge_eval._generate_cell_names("BD2:4")

    assert cells == ["BD2", "BD3", "BD4"]


def test_answer_position_split_preserves_commas_inside_quoted_sheet_names():
    segments = online_judge_eval._split_answer_position(
        "'b2b, sez, de'!A5:V10,Sheet1!A1:B2"
    )

    assert segments == ["'b2b, sez, de'!A5:V10", "Sheet1!A1:B2"]


def test_compare_workbooks_handles_whole_columns_and_quoted_commas(
    tmp_path: Path,
) -> None:
    golden = tmp_path / "golden.xlsx"
    processed = tmp_path / "processed.xlsx"
    workbook = openpyxl.Workbook()
    default = workbook.active
    workbook.remove(default)
    sheet3 = workbook.create_sheet("Sheet3")
    sheet4 = workbook.create_sheet("Sheet4")
    sheet1 = workbook.create_sheet("Sheet1")
    comma_sheet = workbook.create_sheet("b2b, sez, de")
    for sheet in (sheet3, sheet4):
        sheet.append(["a", "b", "c", "d", "e", "f", "g"])
        sheet.append([1, 2, 3, 4, 5, 6, 7])
    comma_sheet.append([None] * 22)
    comma_sheet.append([None] * 22)
    comma_sheet.append([None] * 22)
    comma_sheet.append([None] * 22)
    comma_sheet.append(list(range(1, 23)))
    comma_sheet.append(list(range(23, 45)))
    for row in range(2, 7):
        sheet1.cell(row=row, column=56, value=row)
    workbook.save(golden)
    workbook.save(processed)

    passed, msg = online_judge_eval.compare_workbooks(
        str(golden),
        str(processed),
        "Sheet-Level Manipulation",
        "Sheet3'!A:G,'Sheet4'!A:G,'b2b, sez, de'!A5:V6,'Sheet1'!BD2:6",
    )

    assert passed is True
    assert msg == ""


def test_compare_workbooks_reports_missing_golden_sheet_without_crashing(
    tmp_path: Path,
) -> None:
    golden = tmp_path / "golden.xlsx"
    processed = tmp_path / "processed.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Sheet1"
    workbook.save(golden)
    workbook.save(processed)

    passed, msg = online_judge_eval.compare_workbooks(
        str(golden),
        str(processed),
        "Sheet-Level Manipulation",
        "SANDBOX!A1",
    )

    assert passed is False
    assert msg == "golden worksheet not found: SANDBOX"
