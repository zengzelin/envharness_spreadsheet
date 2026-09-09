# Copyright 2026 The EnvHarness Authors.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Vendored SpreadsheetBench Online-Judge evaluation.

Self-contained copy of the official scorer so the Bridge has no dependency on
the upstream repo layout. Two pieces, both faithful to upstream:

  - `recalc_with_libreoffice(path)` -- the `open_spreadsheet.py` step: drive
    LibreOffice Calc headlessly to recalculate formulas and re-save, so that
    `openpyxl.load_workbook(data_only=True)` reads computed values rather than
    `None`. Required for both the agent's output AND the golden before any
    value comparison.

  - `compare_workbooks(...)` -- the official `evaluation.py` cell comparison:
    for each `Sheet!Range` in `answer_position`, compare every cell value with
    the same rounding / type-normalisation upstream uses. ALL ranges must
    match for the task to pass.

Source: RUCKBReasoning/SpreadsheetBench `evaluation/{evaluation.py,
open_spreadsheet.py}` (value comparison only; the fill/font-colour checks are
commented out upstream and omitted here too).
"""
from __future__ import annotations

import datetime
import os
import re
import shutil
import subprocess
import tempfile

import openpyxl


# ---------------------------------------------------------------------------
# Formula recalculation (LibreOffice headless)
# ---------------------------------------------------------------------------

def find_soffice(explicit: str | None = None) -> str | None:
    for cand in (explicit, "soffice", "libreoffice",
                 "/usr/bin/soffice", "/usr/bin/libreoffice"):
        if not cand:
            continue
        path = shutil.which(cand) if os.path.basename(cand) == cand else cand
        if path and (shutil.which(path) or os.path.isfile(path)):
            return path
    return None


def recalc_with_libreoffice(path: str, soffice_path: str | None = None,
                            timeout: int = 120) -> bool:
    """Open `path` in headless LibreOffice and re-save as xlsx in place so
    formula values get cached. Returns True on success.

    Mirrors `open_spreadsheet.just_open_libreoffice`: convert into a temp dir
    then move the result back over the original.
    """
    soffice = find_soffice(soffice_path)
    if not soffice:
        raise RuntimeError(
            "LibreOffice (soffice) not found. Install it (Linux: "
            "`sudo apt install libreoffice-calc`) or pass "
            "reset_options['soffice_path']."
        )
    path = os.path.abspath(path)
    if not os.path.isfile(path):
        return False
    name = os.path.splitext(os.path.basename(path))[0]
    with tempfile.TemporaryDirectory() as tmp:
        # Per-process LibreOffice user profile: concurrent headless soffice
        # procs sharing one profile lock each other -> random convert failures
        # (the classic headless-concurrency trap). Isolate by pid.
        profile = "file:///tmp/lo_profile_%d" % os.getpid()
        try:
            r = subprocess.run(
                [soffice, "-env:UserInstallation=" + profile,
                 "--headless", "--calc",
                 "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                 "--outdir", tmp, path],
                capture_output=True, text=True, timeout=timeout,
            )
        except subprocess.TimeoutExpired:
            return False
        if r.returncode != 0:
            return False
        converted = os.path.join(tmp, name + ".xlsx")
        if not os.path.isfile(converted):
            return False
        shutil.move(converted, path)
        return True


# ---------------------------------------------------------------------------
# Cell value comparison (official evaluation.py)
# ---------------------------------------------------------------------------

def _datetime_to_float(dt: datetime.datetime) -> float:
    excel_start = datetime.datetime(1899, 12, 30)
    delta = dt - excel_start
    return delta.days + delta.seconds / 86400.0


def _transform_value(v):
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    if isinstance(v, datetime.time):
        return str(v)[:-3]
    if isinstance(v, datetime.datetime):
        return round(_datetime_to_float(v), 0)
    if isinstance(v, str):
        try:
            return round(float(v), 2)
        except ValueError:
            return v
    return v


def _compare_cell_value(v1, v2) -> bool:
    v1 = _transform_value(v1)
    v2 = _transform_value(v2)
    if (v1 == "" and v2 is None) or (v1 is None and v2 == ""):
        return True
    if (v1 == "" and v2 == "") or (v1 is None and v2 is None):
        return True
    if type(v1) != type(v2):
        return False
    return v1 == v2


def _col_num2name(n: int) -> str:
    name = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        name = chr(65 + rem) + name
    return name


def _col_name2num(name: str) -> int:
    num = 0
    for c in name:
        num = num * 26 + (ord(c) - ord("A") + 1)
    return num


_CELL_RE = re.compile(r"^([A-Za-z]*)([0-9]*)$")


def _split_answer_position(answer_position: str) -> list[str]:
    """Split comma-separated Sheet!Range segments, preserving quoted commas."""
    segments: list[str] = []
    start = 0
    in_quote = False
    for index, char in enumerate(answer_position):
        if char == "'":
            # Some SpreadsheetBench rows use Sheet3'!A:G rather than
            # 'Sheet3'!A:G. Treat a quote immediately before ! as closing
            # punctuation, not the start of a quoted sheet name.
            if not in_quote and index + 1 < len(answer_position):
                if answer_position[index + 1] == "!":
                    continue
            in_quote = not in_quote
        elif char == "," and not in_quote:
            segment = answer_position[start:index].strip()
            if segment:
                segments.append(segment)
            start = index + 1
    tail = answer_position[start:].strip()
    if tail:
        segments.append(tail)
    return segments


def _split_cell_ref(cell_ref: str) -> tuple[str, str]:
    match = _CELL_RE.match(cell_ref.strip())
    if not match:
        raise ValueError(f"invalid cell reference: {cell_ref!r}")
    return match.group(1).upper(), match.group(2)


def _parse_cell_range(range_str: str, max_row: int | None = None):
    start_cell, end_cell = range_str.split(":")
    sc, sr = _split_cell_ref(start_cell)
    ec, er = _split_cell_ref(end_cell)
    if sc and sr and not ec and er:
        ec = sc
    if not sr and not er and sc and ec:
        if max_row is None:
            raise ValueError(
                f"whole-column range {range_str!r} needs max_row"
            )
        sr, er = "1", str(max_row)
    if not sc or not sr or not ec or not er:
        raise ValueError(f"invalid cell range: {range_str!r}")
    return (_col_name2num(sc), int(sr)), (_col_name2num(ec), int(er))


def _generate_cell_names(range_str: str, max_row: int | None = None) -> list[str]:
    if ":" not in range_str:
        return [range_str]
    (sc, sr), (ec, er) = _parse_cell_range(range_str, max_row=max_row)
    cols = [_col_num2name(i) for i in range(sc, ec + 1)]
    return [f"{c}{r}" for c in cols for r in range(sr, er + 1)]


def _cell_level_compare(wb_gt, wb_proc, sheet_name: str, cell_range: str):
    if sheet_name not in wb_gt:
        return False, f"golden worksheet not found: {sheet_name}"
    if sheet_name not in wb_proc:
        return False, f"processed worksheet not found: {sheet_name}"
    ws_gt = wb_gt[sheet_name]
    ws_proc = wb_proc[sheet_name]
    max_row = max(ws_gt.max_row, ws_proc.max_row)
    for cell_name in _generate_cell_names(cell_range, max_row=max_row):
        cg = ws_gt[cell_name]
        cp = ws_proc[cell_name]
        if not _compare_cell_value(cg.value, cp.value):
            return False, (f"value diff at {cg.coordinate}: gt={cg.value!r} "
                           f"proc={cp.value!r}")
    return True, ""


def compare_workbooks(gt_file: str, proc_file: str,
                      instruction_type: str, answer_position: str):
    """Return (passed: bool, msg: str). All comma-separated `Sheet!Range`
    segments in `answer_position` must match. Both files are read with
    `data_only=True`, so recalc them first."""
    if not os.path.exists(proc_file):
        return False, "output file does not exist"
    try:
        wb_gt = openpyxl.load_workbook(gt_file, data_only=True)
        wb_proc = openpyxl.load_workbook(proc_file, data_only=True)
    except Exception as e:  # noqa: BLE001
        return False, f"load error: {e}"

    results, msgs = [], []
    for segment in _split_answer_position(answer_position):
        if "!" in segment:
            sheet_name, cell_range = segment.rsplit("!", 1)
        else:
            sheet_name, cell_range = wb_gt.sheetnames[0], segment
        sheet_name = sheet_name.strip().strip("'")
        cell_range = cell_range.strip().strip("'")
        ok, msg = _cell_level_compare(wb_gt, wb_proc, sheet_name, cell_range)
        results.append(ok)
        msgs.append(msg)
    return all(results), "; ".join(m for m in msgs if m)
