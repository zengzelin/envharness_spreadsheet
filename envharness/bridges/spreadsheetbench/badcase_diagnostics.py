"""Bounded, score-neutral diagnostics for failed SpreadsheetBench workbooks."""
from __future__ import annotations

from dataclasses import dataclass
import os
from typing import Any, Iterator

import openpyxl

from . import online_judge_eval


MAX_SERIALIZED_CELL_CHARS = 200


@dataclass(frozen=True)
class DiagnosticLimits:
    max_scan_cells: int = 200_000
    max_examples: int = 20

    def __post_init__(self) -> None:
        if self.max_scan_cells <= 0:
            raise ValueError("max_scan_cells must be positive")
        if self.max_examples <= 0:
            raise ValueError("max_examples must be positive")


class _ScanBudget:
    def __init__(self, maximum: int) -> None:
        self.maximum = maximum
        self.seen: set[tuple[str, int, int]] = set()
        self.truncated = False

    def consume(self, key: tuple[str, int, int]) -> bool:
        if key in self.seen:
            return True
        if len(self.seen) >= self.maximum:
            self.truncated = True
            return False
        self.seen.add(key)
        return True


def _bounded_repr(value: Any) -> str:
    rendered = repr(value)
    if len(rendered) <= MAX_SERIALIZED_CELL_CHARS:
        return rendered
    marker = "...[truncated]"
    return rendered[:MAX_SERIALIZED_CELL_CHARS - len(marker)] + marker


def _color_signature(color: Any) -> tuple[Any, ...] | None:
    if color is None:
        return None
    return (
        getattr(color, "type", None),
        getattr(color, "rgb", None),
        getattr(color, "indexed", None),
        getattr(color, "theme", None),
        getattr(color, "tint", None),
    )


def _side_signature(side: Any) -> tuple[Any, ...] | None:
    if side is None:
        return None
    return (getattr(side, "style", None), _color_signature(side.color))


def _style_signature(cell: Any) -> tuple[Any, ...] | None:
    if cell is None:
        return None
    font = cell.font
    fill = cell.fill
    border = cell.border
    alignment = cell.alignment
    protection = cell.protection
    return (
        cell.number_format,
        (
            font.name, font.sz, font.bold, font.italic, font.vertAlign,
            font.underline, font.strike, _color_signature(font.color),
        ),
        (
            fill.fill_type, _color_signature(fill.fgColor),
            _color_signature(fill.bgColor),
        ),
        (
            _side_signature(border.left), _side_signature(border.right),
            _side_signature(border.top), _side_signature(border.bottom),
            _side_signature(border.diagonal), border.diagonalUp,
            border.diagonalDown,
        ),
        (
            alignment.horizontal, alignment.vertical, alignment.text_rotation,
            alignment.wrap_text, alignment.shrink_to_fit, alignment.indent,
        ),
        (protection.locked, protection.hidden),
    )


def _raw_signature(cell: Any) -> tuple[Any, ...]:
    if cell is None:
        return (None, None, None)
    return (cell.value, cell.data_type, _style_signature(cell))


def _cell(workbook: Any, sheet_name: str, row: int, column: int) -> Any:
    if sheet_name not in workbook.sheetnames:
        return None
    return workbook[sheet_name].cell(row=row, column=column)


def _sheet_max_row(workbooks: tuple[Any, ...], sheet_name: str) -> int:
    return max(
        (
            workbook[sheet_name].max_row
            for workbook in workbooks
            if sheet_name in workbook.sheetnames
        ),
        default=1,
    )


def _sheet_dimensions(workbooks: tuple[Any, ...], sheet_name: str) -> tuple[int, int]:
    rows = []
    columns = []
    for workbook in workbooks:
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        rows.append(max(sheet.max_row, 1))
        columns.append(max(sheet.max_column, 1))
    return max(rows, default=1), max(columns, default=1)


def _answer_ranges(
    answer_position: str,
    *,
    default_sheet: str,
    workbooks: tuple[Any, ...],
) -> Iterator[tuple[str, int, int, int, int]]:
    for segment in online_judge_eval._split_answer_position(answer_position):
        if "!" in segment:
            sheet_name, range_text = segment.rsplit("!", 1)
            sheet_name = sheet_name.strip().strip("'")
        else:
            sheet_name = default_sheet
            range_text = segment
        range_text = range_text.strip().strip("'")
        (min_col, min_row), (max_col, max_row) = (
            online_judge_eval._parse_cell_range(
                range_text,
                max_row=_sheet_max_row(workbooks, sheet_name),
            )
        )
        yield sheet_name, min_col, min_row, max_col, max_row


def _empty_result(status: str) -> dict[str, Any]:
    return {
        "version": 1,
        "status": status,
        "answer_cells_total": 0,
        "answer_cells_scanned": 0,
        "answer_cells_matched": 0,
        "answer_cells_mismatched": 0,
        "answer_match_ratio": 0.0,
        "answer_cells_changed_from_input": 0,
        "answer_required_change_count": 0,
        "answer_range_untouched": False,
        "outside_answer_cells_changed": 0,
        "mutated_sheets": [],
        "formula_cells_expected": 0,
        "formula_cells_output": 0,
        "formula_to_static_count": 0,
        "formula_result_mismatch_count": 0,
        "format_mismatch_count": 0,
        "diagnostic_cells_scanned": 0,
        "diagnostic_truncated": False,
        "first_mismatches": [],
    }


def diagnose_failed_workbook(
    input_path: str,
    golden_path: str,
    output_path: str,
    answer_position: str,
    *,
    limits: DiagnosticLimits,
) -> dict[str, Any]:
    """Compare input/output/golden evidence without changing official grading."""
    if not os.path.isfile(output_path):
        return _empty_result("output_missing")

    workbooks: list[Any] = []
    try:
        input_raw = openpyxl.load_workbook(input_path, data_only=False)
        workbooks.append(input_raw)
        input_values = openpyxl.load_workbook(input_path, data_only=True)
        workbooks.append(input_values)
        golden_raw = openpyxl.load_workbook(golden_path, data_only=False)
        workbooks.append(golden_raw)
        golden_values = openpyxl.load_workbook(golden_path, data_only=True)
        workbooks.append(golden_values)
        output_raw = openpyxl.load_workbook(output_path, data_only=False)
        workbooks.append(output_raw)
        output_values = openpyxl.load_workbook(output_path, data_only=True)
        workbooks.append(output_values)
    except Exception as exc:  # noqa: BLE001
        for workbook in workbooks:
            workbook.close()
        result = _empty_result("load_error")
        result["diagnostic_error_type"] = type(exc).__name__
        return result

    result = _empty_result("complete")
    budget = _ScanBudget(limits.max_scan_cells)
    answer_coordinates: set[tuple[str, int, int]] = set()
    mutated_sheets: set[str] = set()
    try:
        default_sheet = golden_raw.sheetnames[0]
        ranges = list(_answer_ranges(
            answer_position,
            default_sheet=default_sheet,
            workbooks=(input_raw, golden_raw, output_raw),
        ))
        result["answer_cells_total"] = sum(
            (max_col - min_col + 1) * (max_row - min_row + 1)
            for _, min_col, min_row, max_col, max_row in ranges
        )

        answer_budget_exhausted = False
        for sheet_name, min_col, min_row, max_col, max_row in ranges:
            for row in range(min_row, max_row + 1):
                for column in range(min_col, max_col + 1):
                    key = (sheet_name, row, column)
                    if key in answer_coordinates:
                        continue
                    if not budget.consume(key):
                        answer_budget_exhausted = True
                        break
                    answer_coordinates.add(key)
                    result["answer_cells_scanned"] += 1

                    input_value_cell = _cell(input_values, sheet_name, row, column)
                    golden_value_cell = _cell(golden_values, sheet_name, row, column)
                    output_value_cell = _cell(output_values, sheet_name, row, column)
                    input_value = None if input_value_cell is None else input_value_cell.value
                    golden_value = None if golden_value_cell is None else golden_value_cell.value
                    output_value = None if output_value_cell is None else output_value_cell.value

                    matched = online_judge_eval._compare_cell_value(
                        golden_value, output_value
                    )
                    result["answer_cells_matched"] += int(matched)
                    result["answer_cells_mismatched"] += int(not matched)
                    output_changed = not online_judge_eval._compare_cell_value(
                        input_value, output_value
                    )
                    required_change = not online_judge_eval._compare_cell_value(
                        input_value, golden_value
                    )
                    result["answer_cells_changed_from_input"] += int(output_changed)
                    result["answer_required_change_count"] += int(required_change)

                    input_raw_cell = _cell(input_raw, sheet_name, row, column)
                    golden_raw_cell = _cell(golden_raw, sheet_name, row, column)
                    output_raw_cell = _cell(output_raw, sheet_name, row, column)
                    if _raw_signature(input_raw_cell) != _raw_signature(output_raw_cell):
                        mutated_sheets.add(sheet_name)

                    golden_is_formula = bool(
                        golden_raw_cell is not None
                        and golden_raw_cell.data_type == "f"
                    )
                    output_is_formula = bool(
                        output_raw_cell is not None
                        and output_raw_cell.data_type == "f"
                    )
                    result["formula_cells_expected"] += int(golden_is_formula)
                    result["formula_cells_output"] += int(output_is_formula)
                    result["formula_to_static_count"] += int(
                        golden_is_formula and not output_is_formula
                    )
                    result["formula_result_mismatch_count"] += int(
                        golden_is_formula and output_is_formula and not matched
                    )
                    result["format_mismatch_count"] += int(
                        _style_signature(golden_raw_cell)
                        != _style_signature(output_raw_cell)
                    )

                    if not matched and len(result["first_mismatches"]) < limits.max_examples:
                        result["first_mismatches"].append({
                            "sheet": _bounded_repr(sheet_name),
                            "coordinate": _bounded_repr(
                                output_raw_cell.coordinate
                                if output_raw_cell is not None
                                else golden_raw_cell.coordinate
                                if golden_raw_cell is not None
                                else f"R{row}C{column}"
                            ),
                            "initial_value": _bounded_repr(input_value),
                            "output_value": _bounded_repr(output_value),
                            "golden_value": _bounded_repr(golden_value),
                            "output_formula": _bounded_repr(
                                output_raw_cell.value if output_is_formula else None
                            ),
                            "golden_formula": _bounded_repr(
                                golden_raw_cell.value if golden_is_formula else None
                            ),
                            "output_type": _bounded_repr(
                                output_raw_cell.data_type
                                if output_raw_cell is not None else None
                            ),
                            "golden_type": _bounded_repr(
                                golden_raw_cell.data_type
                                if golden_raw_cell is not None else None
                            ),
                        })
                if answer_budget_exhausted:
                    break
            if answer_budget_exhausted:
                break

        scanned = result["answer_cells_scanned"]
        result["answer_match_ratio"] = (
            result["answer_cells_matched"] / scanned if scanned else 0.0
        )
        result["answer_range_untouched"] = bool(
            scanned == result["answer_cells_total"]
            and result["answer_required_change_count"] > 0
            and result["answer_cells_changed_from_input"] == 0
        )

        outside_exhausted = answer_budget_exhausted
        if not outside_exhausted:
            sheet_names = sorted(set(input_raw.sheetnames) | set(output_raw.sheetnames))
            for sheet_name in sheet_names:
                max_row, max_column = _sheet_dimensions(
                    (input_raw, output_raw), sheet_name
                )
                for row in range(1, max_row + 1):
                    for column in range(1, max_column + 1):
                        key = (sheet_name, row, column)
                        if key in answer_coordinates:
                            continue
                        if not budget.consume(key):
                            outside_exhausted = True
                            break
                        if _raw_signature(
                            _cell(input_raw, sheet_name, row, column)
                        ) != _raw_signature(
                            _cell(output_raw, sheet_name, row, column)
                        ):
                            result["outside_answer_cells_changed"] += 1
                            mutated_sheets.add(sheet_name)
                    if outside_exhausted:
                        break
                if outside_exhausted:
                    break

        result["mutated_sheets"] = sorted(mutated_sheets)
        result["diagnostic_cells_scanned"] = len(budget.seen)
        result["diagnostic_truncated"] = bool(budget.truncated)
        return result
    finally:
        for workbook in workbooks:
            workbook.close()
