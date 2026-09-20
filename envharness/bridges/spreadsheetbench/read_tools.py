# Copyright 2026 The EnvHarness Authors.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0

"""Read-only spreadsheet tools used by the SpreadsheetBench bridge."""
from __future__ import annotations

from datetime import date, datetime, time
import json
import math
from pathlib import Path
from typing import Any


NATIVE_READ_TOOLS = frozenset({"list_sheets", "inspect_range", "find_cells"})
NATIVE_WRITE_TOOLS = frozenset({"write_range", "clear_range"})
TOOL_SET_PYTHON = "python"
TOOL_SET_NATIVE_READ = "native_read"
TOOL_SET_NATIVE_BASIC = "native_basic"
VALID_TOOL_SETS = frozenset({
    TOOL_SET_PYTHON, TOOL_SET_NATIVE_READ, TOOL_SET_NATIVE_BASIC
})

MAX_INSPECT_CELLS = 400
MAX_SUMMARY_CELLS = 100_000
MAX_FIND_RESULTS = 1_000
MAX_RESPONSE_CHARS = 10_240
MAX_WORKBOOK_BYTES = 100 * 1024 * 1024


class ReadToolError(ValueError):
    """A user-correctable read-tool argument or workbook error."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


def normalize_tool_set(value: Any) -> str:
    normalized = str(value or TOOL_SET_PYTHON).strip().lower().replace("-", "_")
    if normalized not in VALID_TOOL_SETS:
        choices = ", ".join(sorted(VALID_TOOL_SETS))
        raise ValueError(f"tool_set must be one of {choices}, got {value!r}")
    return normalized


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value)


def _render_payload(payload: dict[str, Any], list_key: str | None = None) -> str:
    text = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    if len(text) <= MAX_RESPONSE_CHARS:
        return text
    if not list_key:
        minimal = {
            "status": payload.get("status", "success"),
            "error": payload.get("error", "response_too_large"),
            "message": str(payload.get("message", ""))[:1000],
            "truncated": True,
        }
        return json.dumps(minimal, ensure_ascii=False, separators=(",", ":"))

    items = payload.get(list_key)
    if not isinstance(items, list):
        return text[:MAX_RESPONSE_CHARS]
    original_count = len(items)
    while items and len(text) > MAX_RESPONSE_CHARS:
        items.pop()
        payload["truncated"] = True
        payload["returned"] = len(items)
        payload["omitted"] = original_count - len(items)
        text = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    if len(text) > MAX_RESPONSE_CHARS:
        payload = {
            "status": payload.get("status", "success"),
            "truncated": True,
            "returned": 0,
            "omitted": original_count,
        }
        text = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    return text


def _require_no_unknown(arguments: dict[str, Any], allowed: set[str]) -> None:
    unknown = sorted(set(arguments) - allowed)
    if unknown:
        raise ReadToolError(
            "invalid_arguments", f"unknown argument(s): {', '.join(unknown)}"
        )


def _open_workbook(path: str, *, data_only: bool = False):
    workbook_path = Path(path)
    if not path or not workbook_path.is_file():
        raise ReadToolError("file_not_found", "current output workbook was not found")
    try:
        size = workbook_path.stat().st_size
    except OSError as exc:
        raise ReadToolError("file_stat_failed", str(exc)) from exc
    if size > MAX_WORKBOOK_BYTES:
        raise ReadToolError(
            "file_too_large",
            f"current output workbook exceeds {MAX_WORKBOOK_BYTES // (1024 * 1024)} MB",
        )
    try:
        import openpyxl

        return openpyxl.load_workbook(
            workbook_path, data_only=data_only, read_only=True, keep_links=False
        )
    except ReadToolError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ReadToolError("workbook_open_failed", str(exc)) from exc


def _split_range(range_text: Any, sheet_name: Any) -> tuple[str | None, str]:
    if not isinstance(range_text, str) or not range_text.strip():
        raise ReadToolError("invalid_range", "range must be a non-empty A1 range")
    token = range_text.strip()
    qualified_sheet: str | None = None
    if "!" in token:
        sheet_token, token = token.rsplit("!", 1)
        sheet_token = sheet_token.strip()
        if len(sheet_token) >= 2 and sheet_token[0] == sheet_token[-1] == "'":
            sheet_token = sheet_token[1:-1].replace("''", "'")
        qualified_sheet = sheet_token
    explicit_sheet = None
    if sheet_name is not None and str(sheet_name).strip():
        explicit_sheet = str(sheet_name).strip()
    if qualified_sheet and explicit_sheet and qualified_sheet != explicit_sheet:
        raise ReadToolError(
            "sheet_mismatch", "sheet_name conflicts with the sheet-qualified range"
        )
    return qualified_sheet or explicit_sheet, token.strip()


def _range_bounds(range_text: str) -> tuple[int, int, int, int]:
    try:
        from openpyxl.utils.cell import range_boundaries

        min_col, min_row, max_col, max_row = range_boundaries(range_text)
    except Exception as exc:  # noqa: BLE001
        raise ReadToolError("invalid_range", "range must be a finite A1 rectangle") from exc
    values = (min_col, min_row, max_col, max_row)
    if any(not isinstance(value, int) or value < 1 for value in values):
        raise ReadToolError("invalid_range", "range must be a finite A1 rectangle")
    if min_col > max_col or min_row > max_row:
        raise ReadToolError("invalid_range", "range bounds are reversed")
    return values


def _resolve_sheet(workbook, sheet_name: str | None):
    if sheet_name is None:
        if not workbook.sheetnames:
            raise ReadToolError("sheet_not_found", "workbook has no worksheets")
        return workbook[workbook.sheetnames[0]]
    if sheet_name not in workbook.sheetnames:
        raise ReadToolError("sheet_not_found", f"worksheet not found: {sheet_name}")
    return workbook[sheet_name]


def _list_sheets(path: str, arguments: dict[str, Any]) -> tuple[dict[str, Any], str]:
    _require_no_unknown(arguments, set())
    workbook = _open_workbook(path)
    try:
        active_index = workbook.index(workbook.active)
        sheets = [
            {
                "name": worksheet.title,
                "index": index,
                "state": worksheet.sheet_state,
                "active": index == active_index,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
            }
            for index, worksheet in enumerate(workbook.worksheets)
        ]
    finally:
        workbook.close()
    payload = {
        "status": "success",
        "sheet_count": len(sheets),
        "sheets": sheets,
        "truncated": False,
    }
    return payload, _render_payload(payload, "sheets")


def _inspect_range(path: str, arguments: dict[str, Any]) -> tuple[dict[str, Any], str]:
    _require_no_unknown(
        arguments, {"range", "sheet_name", "include_details", "mode"}
    )
    sheet_name, range_text = _split_range(
        arguments.get("range"), arguments.get("sheet_name")
    )
    min_col, min_row, max_col, max_row = _range_bounds(range_text)
    cell_count = (max_col - min_col + 1) * (max_row - min_row + 1)
    mode = str(arguments.get("mode", "cells") or "cells").strip().lower()
    if mode not in {"cells", "summary"}:
        raise ReadToolError("invalid_mode", "mode must be 'cells' or 'summary'")
    include_details = arguments.get("include_details", False)
    if not isinstance(include_details, bool):
        raise ReadToolError("invalid_arguments", "include_details must be boolean")
    if mode == "summary" and include_details:
        raise ReadToolError(
            "invalid_mode", "include_details is unavailable in summary mode"
        )
    if mode == "summary" and cell_count > MAX_SUMMARY_CELLS:
        raise ReadToolError(
            "range_too_large",
            f"summary range has {cell_count} cells; maximum is {MAX_SUMMARY_CELLS}",
        )

    workbook = _open_workbook(path)
    try:
        from openpyxl.utils.cell import get_column_letter

        worksheet = _resolve_sheet(workbook, sheet_name)
        if mode == "summary":
            non_empty = formula_cells = error_cells = 0
            for row in worksheet.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_col,
                max_col=max_col,
            ):
                for cell in row:
                    if cell.value is not None:
                        non_empty += 1
                    if cell.data_type == "f" or (
                        isinstance(cell.value, str) and cell.value.startswith("=")
                    ):
                        formula_cells += 1
                    if cell.data_type == "e":
                        error_cells += 1
            payload = {
                "status": "success",
                "sheet": worksheet.title,
                "range": range_text,
                "summary": {
                    "cells": cell_count,
                    "non_empty_cells": non_empty,
                    "empty_cells": cell_count - non_empty,
                    "formula_cells": formula_cells,
                    "error_cells": error_cells,
                    "non_empty_fraction": non_empty / cell_count,
                },
                "truncated": False,
            }
            return payload, _render_payload(payload)

        records: list[dict[str, Any]] = []
        returned = min(cell_count, MAX_INSPECT_CELLS)
        seen = 0
        for row_index, row in enumerate(worksheet.iter_rows(
            min_row=min_row,
            max_row=max_row,
            min_col=min_col,
            max_col=max_col,
        ), start=min_row):
            for column_index, cell in enumerate(row, start=min_col):
                if seen >= returned:
                    break
                value = cell.value
                coordinate = f"{get_column_letter(column_index)}{row_index}"
                record: dict[str, Any] = {"address": coordinate}
                if cell.data_type == "f" or (
                    isinstance(value, str) and value.startswith("=")
                ):
                    record["formula"] = str(value)
                else:
                    record["value"] = _json_value(value)
                if include_details:
                    record["number_format"] = getattr(
                        cell, "number_format", "General"
                    )
                    record["style_id"] = getattr(cell, "style_id", 0)
                    record["data_type"] = getattr(cell, "data_type", "n")
                records.append(record)
                seen += 1
            if seen >= returned:
                break
        payload = {
            "status": "success",
            "sheet": worksheet.title,
            "range": range_text,
            "requested_cells": cell_count,
            "returned": len(records),
            "omitted": cell_count - len(records),
            "cells": records,
            "truncated": cell_count > len(records),
        }
        return payload, _render_payload(payload, "cells")
    finally:
        workbook.close()


def _matches(candidate: Any, query: str, mode: str, case_sensitive: bool) -> bool:
    if candidate is None:
        return False
    candidate_text = str(candidate)
    query_text = query
    if not case_sensitive:
        candidate_text = candidate_text.casefold()
        query_text = query_text.casefold()
    if mode == "equals":
        return candidate_text == query_text
    if mode == "prefix":
        return candidate_text.startswith(query_text)
    return query_text in candidate_text


def _find_cells(path: str, arguments: dict[str, Any]) -> tuple[dict[str, Any], str]:
    _require_no_unknown(
        arguments,
        {
            "query", "sheet_name", "range", "match", "search_in",
            "case_sensitive", "include_values", "return", "return_mode",
            "max_results",
        },
    )
    query = arguments.get("query")
    if not isinstance(query, str) or not query.strip():
        raise ReadToolError("invalid_query", "query must be a non-empty string")
    query = query.strip()
    match_mode = str(arguments.get("match", "contains") or "contains").lower()
    if match_mode not in {"contains", "equals", "prefix"}:
        raise ReadToolError(
            "invalid_match", "match must be contains, equals, or prefix"
        )
    search_in = str(arguments.get("search_in", "values") or "values").lower()
    if search_in not in {"values", "formulas", "both"}:
        raise ReadToolError(
            "invalid_search_in", "search_in must be values, formulas, or both"
        )
    case_sensitive = arguments.get("case_sensitive", False)
    include_values = arguments.get("include_values", False)
    if not isinstance(case_sensitive, bool) or not isinstance(include_values, bool):
        raise ReadToolError(
            "invalid_arguments", "case_sensitive and include_values must be boolean"
        )
    return_mode = arguments.get("return_mode", arguments.get("return", "first"))
    return_mode = str(return_mode or "first").lower()
    if return_mode not in {"first", "all"}:
        raise ReadToolError("invalid_return_mode", "return mode must be first or all")
    try:
        max_results = int(arguments.get("max_results", 20))
    except (TypeError, ValueError) as exc:
        raise ReadToolError("invalid_max_results", "max_results must be an integer") from exc
    max_results = min(max(max_results, 1), MAX_FIND_RESULTS)
    if return_mode == "first":
        max_results = 1

    range_argument = arguments.get("range")
    sheet_argument = arguments.get("sheet_name")
    if range_argument:
        qualified_sheet, range_text = _split_range(range_argument, sheet_argument)
        bounds = _range_bounds(range_text)
    else:
        qualified_sheet = str(sheet_argument).strip() if sheet_argument else None
        range_text = ""
        bounds = None

    workbook = _open_workbook(path)
    values_workbook = (
        _open_workbook(path, data_only=True)
        if search_in in {"values", "both"}
        else None
    )
    try:
        if qualified_sheet:
            worksheets = [_resolve_sheet(workbook, qualified_sheet)]
        else:
            worksheets = list(workbook.worksheets)
        results: list[dict[str, Any]] = []
        scanned = 0
        for worksheet in worksheets:
            values_worksheet = (
                values_workbook[worksheet.title] if values_workbook else None
            )
            if bounds:
                min_col, min_row, max_col, max_row = bounds
            else:
                min_col, min_row = 1, 1
                max_col, max_row = worksheet.max_column, worksheet.max_row
            raw_rows = worksheet.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_col,
                max_col=max_col,
            )
            if values_worksheet:
                value_rows = values_worksheet.iter_rows(
                    min_row=min_row,
                    max_row=max_row,
                    min_col=min_col,
                    max_col=max_col,
                )
                row_pairs = zip(raw_rows, value_rows)
            else:
                row_pairs = ((row, row) for row in raw_rows)
            for row, value_row in row_pairs:
                for cell, value_cell in zip(row, value_row):
                    scanned += 1
                    raw = cell.value
                    is_formula = cell.data_type == "f" or (
                        isinstance(raw, str) and raw.startswith("=")
                    )
                    candidates: list[tuple[str, Any]] = []
                    visible_value = value_cell.value if is_formula else raw
                    if search_in in {"values", "both"}:
                        candidates.append(("value", visible_value))
                    if search_in in {"formulas", "both"} and is_formula:
                        candidates.append(("formula", raw))
                    matched_field = next(
                        (
                            field
                            for field, candidate in candidates
                            if _matches(candidate, query, match_mode, case_sensitive)
                        ),
                        None,
                    )
                    if matched_field is None:
                        continue
                    result: dict[str, Any] = {
                        "sheet": worksheet.title,
                        "address": cell.coordinate,
                    }
                    if include_values or matched_field == "formula":
                        matched_value = raw if matched_field == "formula" else visible_value
                        result[matched_field] = _json_value(matched_value)
                    results.append(result)
                    if len(results) >= max_results:
                        break
                if len(results) >= max_results:
                    break
            if len(results) >= max_results:
                break
        payload = {
            "status": "success",
            "query": query,
            "match": match_mode,
            "search_in": search_in,
            "returned": len(results),
            "scanned_cells": scanned,
            "matches": results,
            "truncated": len(results) >= max_results and return_mode == "all",
        }
        if range_text:
            payload["range"] = range_text
        return payload, _render_payload(payload, "matches")
    finally:
        workbook.close()
        if values_workbook is not None:
            values_workbook.close()


def execute_read_tool(
    name: str, path: str, arguments: dict[str, Any]
) -> tuple[dict[str, Any], str]:
    """Execute one read tool and return both structured and text results."""
    if not isinstance(arguments, dict):
        raise ReadToolError("invalid_arguments", "arguments must be an object")
    if name == "list_sheets":
        return _list_sheets(path, arguments)
    if name == "inspect_range":
        return _inspect_range(path, arguments)
    if name == "find_cells":
        return _find_cells(path, arguments)
    raise ReadToolError("unknown_tool", f"unknown read tool: {name}")


def error_payload(error: ReadToolError) -> tuple[dict[str, Any], str]:
    payload = {"status": "error", "error": error.code, "message": error.message}
    return payload, _render_payload(payload)
