"""Bounded, non-mutating LibreOffice recalculation for SpreadsheetBench."""
from __future__ import annotations

import os
from pathlib import Path
import shutil
import tempfile
import time
from typing import Any

from . import online_judge_eval
from .read_tools import (
    MAX_WORKBOOK_BYTES,
    ReadToolError,
    _json_value,
    _range_bounds,
    _render_payload,
    _require_no_unknown,
    _resolve_sheet,
    _split_range,
)
from .write_tools import _workbook_lock


MAX_RECALC_RANGES = 10
MAX_RECALC_CELLS = 400
DEFAULT_RECALC_TIMEOUT = 120
_FORMULA_ERRORS = frozenset({
    "#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A",
    "#GETTING_DATA", "#SPILL!", "#CALC!", "#FIELD!", "#BLOCKED!",
    "#UNKNOWN!", "#CONNECT!",
})


def execute_recalculate_and_read(
    path: str,
    arguments: dict[str, Any],
    *,
    soffice_path: str | None = None,
    timeout: int = DEFAULT_RECALC_TIMEOUT,
) -> tuple[dict[str, Any], str]:
    """Recalculate a temporary copy and return cached values from finite ranges."""
    _require_no_unknown(arguments, {"cell_ranges"})
    ranges = arguments.get("cell_ranges")
    if not isinstance(ranges, list) or not ranges:
        raise ReadToolError(
            "invalid_ranges", "cell_ranges must be a non-empty list of A1 ranges"
        )
    if len(ranges) > MAX_RECALC_RANGES:
        raise ReadToolError(
            "too_many_ranges",
            f"cell_ranges may contain at most {MAX_RECALC_RANGES} ranges",
        )

    parsed: list[tuple[str | None, str, tuple[int, int, int, int]]] = []
    total_cells = 0
    for item in ranges:
        sheet_name, range_text = _split_range(item, None)
        bounds = _range_bounds(range_text)
        min_col, min_row, max_col, max_row = bounds
        total_cells += (max_col - min_col + 1) * (max_row - min_row + 1)
        parsed.append((sheet_name, range_text, bounds))
    if total_cells > MAX_RECALC_CELLS:
        raise ReadToolError(
            "range_too_large",
            f"requested ranges contain {total_cells} cells; maximum is {MAX_RECALC_CELLS}",
        )

    source = Path(path)
    if not source.is_file():
        raise ReadToolError("file_not_found", "current output workbook was not found")
    if source.stat().st_size > MAX_WORKBOOK_BYTES:
        raise ReadToolError("file_too_large", "current output workbook exceeds 100 MB")

    started = time.monotonic()
    with tempfile.TemporaryDirectory(prefix="sb_recalc_") as temporary_dir:
        temporary_path = Path(temporary_dir) / source.name
        with _workbook_lock(path):
            shutil.copyfile(source, temporary_path)
        try:
            ok = online_judge_eval.recalc_with_libreoffice(
                str(temporary_path), soffice_path, timeout=timeout
            )
        except RuntimeError as exc:
            raise ReadToolError("recalc_unavailable", str(exc)) from exc
        if not ok:
            raise ReadToolError(
                "recalc_failed", "LibreOffice failed or timed out while recalculating"
            )

        try:
            import openpyxl

            values_book = openpyxl.load_workbook(
                temporary_path, data_only=True, read_only=True, keep_links=False
            )
            formula_book = openpyxl.load_workbook(
                temporary_path, data_only=False, read_only=True, keep_links=False
            )
        except Exception as exc:  # noqa: BLE001
            raise ReadToolError("workbook_open_failed", str(exc)) from exc

        results: list[dict[str, Any]] = []
        formula_count = 0
        error_count = 0
        try:
            for sheet_name, range_text, bounds in parsed:
                value_sheet = _resolve_sheet(values_book, sheet_name)
                formula_sheet = _resolve_sheet(formula_book, value_sheet.title)
                min_col, min_row, max_col, max_row = bounds
                matrix: list[list[Any]] = []
                errors: list[dict[str, Any]] = []
                for row in range(min_row, max_row + 1):
                    output_row: list[Any] = []
                    for col in range(min_col, max_col + 1):
                        value = value_sheet.cell(row, col).value
                        formula = formula_sheet.cell(row, col).value
                        if isinstance(formula, str) and formula.startswith("="):
                            formula_count += 1
                        if isinstance(value, str) and value.upper() in _FORMULA_ERRORS:
                            error_count += 1
                            errors.append({
                                "address": value_sheet.cell(row, col).coordinate,
                                "value": value,
                            })
                        output_row.append(_json_value(value))
                    matrix.append(output_row)
                results.append({
                    "sheet": value_sheet.title,
                    "range": range_text,
                    "values": matrix,
                    "errors": errors,
                })
        finally:
            values_book.close()
            formula_book.close()

    payload = {
        "status": "success",
        "updated": False,
        "values_source": "libreoffice_temporary_copy",
        "requested_cells": total_cells,
        "formula_cells": formula_count,
        "formula_error_cells": error_count,
        "elapsed_ms": round((time.monotonic() - started) * 1000, 1),
        "ranges": results,
    }
    return payload, _render_payload(payload, "ranges")
