# Copyright 2026 The EnvHarness Authors.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.

"""Bounded, atomic workbook mutation tools for SpreadsheetBench."""
from __future__ import annotations

from contextlib import contextmanager
import os
from pathlib import Path
import tempfile
import time
from typing import Any, Iterator

from envharness.bridges.spreadsheetbench.read_tools import (
    MAX_WORKBOOK_BYTES,
    ReadToolError,
    _json_value,
    _range_bounds,
    _render_payload,
    _require_no_unknown,
    _split_range,
)


MAX_WRITE_CELLS = 50_000
MAX_FILL_CELLS = 300_000
MAX_CELL_CHARS = 8_192
LOCK_TIMEOUT_SECONDS = 30.0


@contextmanager
def _workbook_lock(path: str) -> Iterator[None]:
    import fcntl

    lock_path = path + ".lock"
    handle = open(lock_path, "a+b")
    deadline = time.monotonic() + LOCK_TIMEOUT_SECONDS
    try:
        while True:
            try:
                fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
                break
            except BlockingIOError:
                if time.monotonic() >= deadline:
                    raise ReadToolError(
                        "lock_timeout", "timed out waiting for workbook lock"
                    )
                time.sleep(0.05)
        yield
    finally:
        try:
            fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
        finally:
            handle.close()


def _open_for_edit(path: str):
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise ReadToolError("file_not_found", "current output workbook was not found")
    if workbook_path.stat().st_size > MAX_WORKBOOK_BYTES:
        raise ReadToolError("file_too_large", "current output workbook exceeds 100 MB")
    try:
        import openpyxl

        return openpyxl.load_workbook(workbook_path, data_only=False)
    except Exception as exc:  # noqa: BLE001
        raise ReadToolError("workbook_open_failed", str(exc)) from exc


def _atomic_save(workbook: Any, path: str) -> None:
    destination = Path(path)
    descriptor, temporary = tempfile.mkstemp(
        prefix=f".{destination.stem}.", suffix=destination.suffix,
        dir=str(destination.parent),
    )
    os.close(descriptor)
    try:
        workbook.save(temporary)
        import openpyxl

        check = openpyxl.load_workbook(temporary, read_only=True, data_only=False)
        check.close()
        os.replace(temporary, destination)
    except ReadToolError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ReadToolError("workbook_save_failed", str(exc)) from exc
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)


def _resolve_write_sheet(workbook: Any, sheet_name: str | None, *, create: bool):
    if sheet_name is None:
        if not workbook.sheetnames:
            return workbook.create_sheet("Sheet1")
        return workbook[workbook.sheetnames[0]]
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    if create:
        return workbook.create_sheet(sheet_name)
    raise ReadToolError("sheet_not_found", f"worksheet not found: {sheet_name}")


def _normalize_data(data: Any) -> tuple[list[list[Any]], bool]:
    scalar_types = (str, int, float, bool)
    if isinstance(data, scalar_types):
        return [[data]], True
    if not isinstance(data, list) or not data:
        raise ReadToolError(
            "invalid_data", "data must be a scalar, non-empty row, or 2D array"
        )
    if all(not isinstance(item, list) for item in data):
        return [data], False
    if not all(isinstance(item, list) for item in data):
        raise ReadToolError("invalid_data", "data cannot mix rows and scalars")
    if not data[0] or any(len(row) != len(data[0]) for row in data):
        raise ReadToolError("invalid_data", "2D data must be non-empty and rectangular")
    return data, False


def _validate_static_value(value: Any) -> None:
    if value is None or isinstance(value, (bool, int, float)):
        return
    if not isinstance(value, str):
        raise ReadToolError(
            "invalid_cell_value", "cell values must be null, string, number, or boolean"
        )
    if len(value) > MAX_CELL_CHARS:
        raise ReadToolError(
            "cell_value_too_large", f"cell strings may not exceed {MAX_CELL_CHARS} characters"
        )
    if value.lstrip().startswith("="):
        raise ReadToolError(
            "formula_not_allowed", "write_range accepts static values only"
        )


def _validate_formula_syntax(formula: str) -> None:
    """Reject structural formula errors that openpyxl otherwise saves silently."""
    delimiter_pairs = {"(": ")", "{": "}"}
    closing_delimiters = {value: key for key, value in delimiter_pairs.items()}
    stack: list[tuple[str, int]] = []
    quote: str | None = None
    quote_start = 0
    index = 0
    while index < len(formula):
        character = formula[index]
        if quote is not None:
            if character == quote:
                if index + 1 < len(formula) and formula[index + 1] == quote:
                    index += 2
                    continue
                quote = None
            index += 1
            continue
        if character in {'"', "'"}:
            quote = character
            quote_start = index
        elif character in delimiter_pairs:
            stack.append((character, index))
        elif character in closing_delimiters:
            expected = closing_delimiters[character]
            if not stack:
                raise ReadToolError(
                    "invalid_formula",
                    f"formula syntax invalid at position {index + 1}: "
                    f"unexpected {character!r}",
                )
            opening, opening_index = stack[-1]
            if opening != expected:
                raise ReadToolError(
                    "invalid_formula",
                    f"formula syntax invalid at position {index + 1}: expected "
                    f"{delimiter_pairs[opening]!r} to close {opening!r} at "
                    f"position {opening_index + 1}",
                )
            stack.pop()
        index += 1

    if quote is not None:
        quote_name = "double quote" if quote == '"' else "single quote"
        raise ReadToolError(
            "invalid_formula",
            f"formula syntax invalid at position {quote_start + 1}: "
            f"unclosed {quote_name}",
        )
    if stack:
        opening, opening_index = stack[-1]
        raise ReadToolError(
            "invalid_formula",
            f"formula syntax invalid at position {opening_index + 1}: "
            f"unclosed {opening!r}",
        )

    try:
        from openpyxl.formula import Tokenizer
        from openpyxl.formula.tokenizer import TokenizerError

        tokens = [
            token for token in Tokenizer(formula).items
            if token.type != "WHITE-SPACE"
        ]
    except (IndexError, TokenizerError, ValueError) as exc:
        raise ReadToolError(
            "invalid_formula", f"formula syntax invalid: {exc}"
        ) from None
    if not tokens:
        raise ReadToolError("invalid_formula", "formula has no expression")

    for token_index, token in enumerate(tokens):
        if token.type not in {"OPERATOR-INFIX", "OPERATOR-PREFIX"}:
            continue
        previous = tokens[token_index - 1] if token_index else None
        following = (
            tokens[token_index + 1]
            if token_index + 1 < len(tokens)
            else None
        )
        if token.type == "OPERATOR-INFIX" and (
            previous is None
            or following is None
            or previous.type in {"OPERATOR-INFIX", "OPERATOR-PREFIX", "SEP"}
            or previous.subtype == "OPEN"
            or following.type in {"OPERATOR-INFIX", "OPERATOR-POSTFIX", "SEP"}
            or following.subtype == "CLOSE"
        ):
            raise ReadToolError(
                "invalid_formula",
                f"formula syntax invalid: operator {token.value!r} "
                "is missing an operand",
            )
        if token.type == "OPERATOR-PREFIX" and (
            following is None
            or following.type == "SEP"
            or following.subtype == "CLOSE"
        ):
            raise ReadToolError(
                "invalid_formula",
                f"formula syntax invalid: operator {token.value!r} "
                "is missing an operand",
            )


def _write_range(path: str, arguments: dict[str, Any]) -> tuple[dict[str, Any], str]:
    _require_no_unknown(arguments, {"range", "sheet_name", "data"})
    if "data" not in arguments:
        raise ReadToolError("missing_data", "write_range requires data")
    sheet_name, range_text = _split_range(
        arguments.get("range"), arguments.get("sheet_name")
    )
    min_col, min_row, max_col, max_row = _range_bounds(range_text)
    rows, scalar = _normalize_data(arguments["data"])
    for row in rows:
        for value in row:
            _validate_static_value(value)

    target_rows = max_row - min_row + 1
    target_cols = max_col - min_col + 1
    if scalar:
        rows = [[rows[0][0] for _ in range(target_cols)] for _ in range(target_rows)]
    elif len(rows) == 1 and target_cols == 1 and len(rows[0]) == target_rows:
        rows = [[value] for value in rows[0]]
    elif target_rows == target_cols == 1:
        target_rows, target_cols = len(rows), len(rows[0])
        max_row = min_row + target_rows - 1
        max_col = min_col + target_cols - 1
    elif len(rows) != target_rows or len(rows[0]) != target_cols:
        raise ReadToolError(
            "shape_mismatch",
            f"data shape {len(rows)}x{len(rows[0])} does not match "
            f"target {target_rows}x{target_cols}",
        )
    cell_count = target_rows * target_cols
    if cell_count > MAX_WRITE_CELLS:
        raise ReadToolError(
            "range_too_large", f"write has {cell_count} cells; maximum is {MAX_WRITE_CELLS}"
        )

    with _workbook_lock(path):
        workbook = _open_for_edit(path)
        try:
            worksheet = _resolve_write_sheet(workbook, sheet_name, create=True)
            changed = 0
            samples: list[dict[str, Any]] = []
            for row_offset, row in enumerate(rows):
                for col_offset, value in enumerate(row):
                    if value is None:
                        continue
                    cell = worksheet.cell(min_row + row_offset, min_col + col_offset)
                    cell.value = value
                    changed += 1
                    if len(samples) < 10:
                        samples.append({
                            "address": cell.coordinate,
                            "value": _json_value(value),
                        })
            _atomic_save(workbook, path)
            resolved_sheet = worksheet.title
        finally:
            workbook.close()
    from openpyxl.utils.cell import get_column_letter

    resolved_range = (
        f"{get_column_letter(min_col)}{min_row}:"
        f"{get_column_letter(max_col)}{max_row}"
    )
    payload = {
        "status": "success", "sheet": resolved_sheet, "range": resolved_range,
        "requested_cells": cell_count, "changed_cells": changed,
        "samples": samples,
    }
    return payload, _render_payload(payload, "samples")


def _clear_range(path: str, arguments: dict[str, Any]) -> tuple[dict[str, Any], str]:
    _require_no_unknown(arguments, {"range", "sheet_name"})
    sheet_name, range_text = _split_range(
        arguments.get("range"), arguments.get("sheet_name")
    )
    min_col, min_row, max_col, max_row = _range_bounds(range_text)
    cell_count = (max_row - min_row + 1) * (max_col - min_col + 1)
    if cell_count > MAX_WRITE_CELLS:
        raise ReadToolError(
            "range_too_large", f"clear has {cell_count} cells; maximum is {MAX_WRITE_CELLS}"
        )

    with _workbook_lock(path):
        workbook = _open_for_edit(path)
        try:
            worksheet = _resolve_write_sheet(workbook, sheet_name, create=False)
            changed = 0
            samples: list[str] = []
            merged_anchors: set[tuple[int, int]] = set()
            for merged in worksheet.merged_cells.ranges:
                if not (
                    merged.max_col < min_col or merged.min_col > max_col
                    or merged.max_row < min_row or merged.min_row > max_row
                ):
                    merged_anchors.add((merged.min_row, merged.min_col))
            coordinates = {
                (row, col)
                for row in range(min_row, max_row + 1)
                for col in range(min_col, max_col + 1)
            } | merged_anchors
            from openpyxl.cell.cell import MergedCell

            for row, col in sorted(coordinates):
                cell = worksheet.cell(row, col)
                if isinstance(cell, MergedCell) or cell.value is None:
                    continue
                cell.value = None
                changed += 1
                if len(samples) < 10:
                    samples.append(cell.coordinate)
            _atomic_save(workbook, path)
            resolved_sheet = worksheet.title
        finally:
            workbook.close()
    payload = {
        "status": "success", "sheet": resolved_sheet, "range": range_text,
        "requested_cells": cell_count, "changed_cells": changed,
        "samples": samples,
    }
    return payload, _render_payload(payload, "samples")


def _fill_formula(path: str, arguments: dict[str, Any]) -> tuple[dict[str, Any], str]:
    _require_no_unknown(
        arguments,
        {"start_cell", "formula_template", "sheet_name", "end_row", "end_col"},
    )
    if "start_cell" not in arguments or "formula_template" not in arguments:
        raise ReadToolError(
            "missing_arguments", "fill_formula requires start_cell and formula_template"
        )
    sheet_name, start_text = _split_range(
        arguments["start_cell"], arguments.get("sheet_name")
    )
    min_col, min_row, max_col, max_row = _range_bounds(start_text)
    if min_col != max_col or min_row != max_row:
        raise ReadToolError("invalid_start_cell", "start_cell must be one A1 cell")

    formula = arguments["formula_template"]
    if not isinstance(formula, str) or not formula.startswith("="):
        raise ReadToolError(
            "invalid_formula", "formula_template must be an Excel formula starting with '='"
        )
    if len(formula) > MAX_CELL_CHARS:
        raise ReadToolError(
            "formula_too_large",
            f"formula_template may not exceed {MAX_CELL_CHARS} characters",
        )
    _validate_formula_syntax(formula)

    end_row = arguments.get("end_row", min_row)
    if isinstance(end_row, bool) or not isinstance(end_row, int) or end_row < min_row:
        raise ReadToolError(
            "invalid_end_row", "end_row must be an integer at or below start_cell"
        )
    end_col_value = arguments.get("end_col")
    if end_col_value is None or end_col_value == "":
        end_col = min_col
    else:
        if not isinstance(end_col_value, str):
            raise ReadToolError("invalid_end_col", "end_col must be a column label")
        try:
            from openpyxl.utils.cell import column_index_from_string

            end_col = column_index_from_string(end_col_value.strip().replace("$", ""))
        except Exception as exc:  # noqa: BLE001
            raise ReadToolError("invalid_end_col", "end_col must be a column label") from exc
        if end_col < min_col:
            raise ReadToolError(
                "invalid_end_col", "end_col must be at or to the right of start_cell"
            )

    cell_count = (end_row - min_row + 1) * (end_col - min_col + 1)
    if cell_count > MAX_FILL_CELLS:
        raise ReadToolError(
            "range_too_large",
            f"formula fill has {cell_count} cells; maximum is {MAX_FILL_CELLS}",
        )

    with _workbook_lock(path):
        workbook = _open_for_edit(path)
        try:
            worksheet = _resolve_write_sheet(workbook, sheet_name, create=False)
            from openpyxl.cell.cell import MergedCell
            from openpyxl.formula.translate import Translator

            origin = worksheet.cell(min_row, min_col).coordinate
            changed = 0
            samples: list[dict[str, Any]] = []
            for row in range(min_row, end_row + 1):
                for col in range(min_col, end_col + 1):
                    cell = worksheet.cell(row, col)
                    if isinstance(cell, MergedCell):
                        continue
                    translated = Translator(formula, origin=origin).translate_formula(
                        cell.coordinate
                    )
                    cell.value = translated
                    changed += 1
                    if len(samples) < 10:
                        samples.append({
                            "address": cell.coordinate,
                            "formula": translated,
                        })
            _atomic_save(workbook, path)
            resolved_sheet = worksheet.title
        finally:
            workbook.close()

    from openpyxl.utils.cell import get_column_letter

    resolved_range = (
        f"{get_column_letter(min_col)}{min_row}:"
        f"{get_column_letter(end_col)}{end_row}"
    )
    payload = {
        "status": "success",
        "sheet": resolved_sheet,
        "range": resolved_range,
        "requested_cells": cell_count,
        "changed_cells": changed,
        "samples": samples,
    }
    return payload, _render_payload(payload, "samples")


def execute_write_tool(
    name: str, path: str, arguments: dict[str, Any]
) -> tuple[dict[str, Any], str]:
    try:
        if not isinstance(arguments, dict):
            raise ReadToolError("invalid_arguments", "arguments must be an object")
        if name == "write_range":
            return _write_range(path, arguments)
        if name == "clear_range":
            return _clear_range(path, arguments)
        if name == "fill_formula":
            return _fill_formula(path, arguments)
        raise ReadToolError("unknown_tool", f"unknown write tool: {name}")
    except ReadToolError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ReadToolError("tool_execution_failed", str(exc)) from exc
