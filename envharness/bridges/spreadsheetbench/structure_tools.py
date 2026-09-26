"""Atomic formatting and worksheet structure tools for SpreadsheetBench."""
from __future__ import annotations

from copy import copy
import re
from typing import Any

from .read_tools import (
    ReadToolError,
    _range_bounds,
    _render_payload,
    _require_no_unknown,
    _split_range,
)
from .write_tools import (
    _atomic_save,
    _open_for_edit,
    _resolve_write_sheet,
    _workbook_lock,
)


_COLOR_RE = re.compile(r"^#?(?:[0-9A-Fa-f]{6}|[0-9A-Fa-f]{8})$")
_SHEET_INVALID_RE = re.compile(r"[\\/*?:\[\]]")
_ROW_RE = re.compile(r"^([1-9]\d*)(?::([1-9]\d*))?$")
_COL_RE = re.compile(r"^([A-Za-z]{1,3})(?::([A-Za-z]{1,3}))?$")


def _color(value: Any, field: str) -> str:
    if not isinstance(value, str) or not _COLOR_RE.fullmatch(value.strip()):
        raise ReadToolError("invalid_color", f"{field} must be RRGGBB or AARRGGBB")
    token = value.strip().lstrip("#").upper()
    return "FF" + token if len(token) == 6 else token


def _format_range(path: str, arguments: dict[str, Any]) -> tuple[dict[str, Any], str]:
    allowed = {
        "range", "sheet_name", "font", "fill", "alignment", "border",
        "number_format", "row_height", "column_width",
    }
    _require_no_unknown(arguments, allowed)
    sheet_name, range_text = _split_range(
        arguments.get("range"), arguments.get("sheet_name")
    )
    min_col, min_row, max_col, max_row = _range_bounds(range_text)
    options = {key: arguments.get(key) for key in allowed - {"range", "sheet_name"}}
    if all(value is None for value in options.values()):
        raise ReadToolError("no_format_options", "provide at least one format option")

    font_options = options["font"] or {}
    fill_options = options["fill"] or {}
    alignment_options = options["alignment"] or {}
    border_options = options["border"] or {}
    for name, value in (
        ("font", font_options), ("fill", fill_options),
        ("alignment", alignment_options), ("border", border_options),
    ):
        if not isinstance(value, dict):
            raise ReadToolError("invalid_arguments", f"{name} must be an object")

    unknown_font = set(font_options) - {
        "bold", "italic", "underline", "size", "name", "color"
    }
    unknown_fill = set(fill_options) - {"color"}
    unknown_alignment = set(alignment_options) - {
        "horizontal", "vertical", "text_rotation", "wrap_text",
        "shrink_to_fit", "indent",
    }
    unknown_border = set(border_options) - {"style", "color"}
    for label, unknown in (
        ("font", unknown_font), ("fill", unknown_fill),
        ("alignment", unknown_alignment), ("border", unknown_border),
    ):
        if unknown:
            raise ReadToolError(
                "invalid_arguments",
                f"unknown {label} option(s): " + ", ".join(sorted(unknown)),
            )
    font_color = (
        _color(font_options["color"], "font.color")
        if "color" in font_options else None
    )
    fill_color = (
        _color(fill_options.get("color"), "fill.color")
        if fill_options else None
    )
    border_color = (
        _color(border_options.get("color", "000000"), "border.color")
        if border_options else None
    )
    number_format = options["number_format"]
    if number_format is not None and not isinstance(number_format, str):
        raise ReadToolError(
            "invalid_number_format", "number_format must be a string"
        )
    row_height = None
    if options["row_height"] is not None:
        try:
            row_height = float(options["row_height"])
        except (TypeError, ValueError):
            raise ReadToolError(
                "invalid_row_height", "row_height must be numeric"
            ) from None
        if not 0 < row_height <= 409:
            raise ReadToolError("invalid_row_height", "row_height must be 0-409")
    column_width = None
    if options["column_width"] is not None:
        try:
            column_width = float(options["column_width"])
        except (TypeError, ValueError):
            raise ReadToolError(
                "invalid_column_width", "column_width must be numeric"
            ) from None
        if not 0 < column_width <= 255:
            raise ReadToolError("invalid_column_width", "column_width must be 0-255")

    with _workbook_lock(path):
        workbook = _open_for_edit(path)
        try:
            worksheet = _resolve_write_sheet(workbook, sheet_name, create=False)
            from openpyxl.styles import Border, PatternFill, Side
            from openpyxl.utils import get_column_letter

            for row in worksheet.iter_rows(
                min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
            ):
                for cell in row:
                    if font_options:
                        font = copy(cell.font)
                        for key in ("bold", "italic", "underline", "size", "name"):
                            if key in font_options:
                                setattr(font, key, font_options[key])
                        if font_color is not None:
                            font.color = font_color
                        cell.font = font
                    if fill_options:
                        cell.fill = PatternFill(fill_type="solid", fgColor=fill_color)
                    if alignment_options:
                        alignment = copy(cell.alignment)
                        for key, value in alignment_options.items():
                            setattr(alignment, key, value)
                        cell.alignment = alignment
                    if border_options:
                        style = border_options.get("style", "thin")
                        side = Side(style=style, color=border_color)
                        cell.border = Border(left=side, right=side, top=side, bottom=side)
                    if number_format is not None:
                        cell.number_format = number_format
            if row_height is not None:
                for row in range(min_row, max_row + 1):
                    worksheet.row_dimensions[row].height = row_height
            if column_width is not None:
                for col in range(min_col, max_col + 1):
                    worksheet.column_dimensions[get_column_letter(col)].width = column_width
            _atomic_save(workbook, path)
            resolved_sheet = worksheet.title
        finally:
            workbook.close()
    payload = {
        "status": "success", "sheet": resolved_sheet, "range": range_text,
        "formatted_cells": (max_row - min_row + 1) * (max_col - min_col + 1),
    }
    return payload, _render_payload(payload)


def _tokens(value: Any, field: str) -> list[str]:
    if isinstance(value, str):
        value = [value]
    if not isinstance(value, list) or not value or not all(
        isinstance(item, str) and item.strip() for item in value
    ):
        raise ReadToolError("invalid_range", f"{field} must be a non-empty range list")
    return [item.strip() for item in value]


def _delete_axis(
    path: str, arguments: dict[str, Any], *, axis: str
) -> tuple[dict[str, Any], str]:
    field = "rows" if axis == "rows" else "columns"
    _require_no_unknown(arguments, {field, "sheet_name"})
    parsed: list[tuple[int, int]] = []
    for token in _tokens(arguments.get(field), field):
        match = (_ROW_RE if axis == "rows" else _COL_RE).fullmatch(token)
        if match is None:
            raise ReadToolError("invalid_range", f"invalid {axis} range: {token}")
        if axis == "rows":
            start, end = int(match.group(1)), int(match.group(2) or match.group(1))
        else:
            from openpyxl.utils import column_index_from_string
            start = column_index_from_string(match.group(1))
            end = column_index_from_string(match.group(2) or match.group(1))
        if end < start:
            start, end = end, start
        parsed.append((start, end))
    parsed.sort()
    merged_intervals: list[tuple[int, int]] = []
    for start, end in parsed:
        if merged_intervals and start <= merged_intervals[-1][1] + 1:
            previous_start, previous_end = merged_intervals[-1]
            merged_intervals[-1] = (previous_start, max(previous_end, end))
        else:
            merged_intervals.append((start, end))
    parsed = list(reversed(merged_intervals))

    with _workbook_lock(path):
        workbook = _open_for_edit(path)
        try:
            worksheet = _resolve_write_sheet(
                workbook, str(arguments.get("sheet_name") or "").strip() or None,
                create=False,
            )
            axis_limit = worksheet.max_row if axis == "rows" else worksheet.max_column
            if any(end > axis_limit for _, end in parsed):
                raise ReadToolError(
                    "range_out_of_bounds",
                    f"cannot delete {axis} beyond used worksheet boundary {axis_limit}",
                )
            for start, end in parsed:
                for merged in worksheet.merged_cells.ranges:
                    hit = (
                        merged.min_row <= end and merged.max_row >= start
                        if axis == "rows"
                        else merged.min_col <= end and merged.max_col >= start
                    )
                    if hit:
                        raise ReadToolError(
                            "merged_range_intersection",
                            f"cannot delete {axis} intersecting merged range {merged}",
                        )
                amount = end - start + 1
                if axis == "rows":
                    worksheet.delete_rows(start, amount)
                else:
                    worksheet.delete_cols(start, amount)
            _atomic_save(workbook, path)
            resolved_sheet = worksheet.title
        finally:
            workbook.close()
    payload = {
        "status": "success", "sheet": resolved_sheet, "axis": axis,
        "deleted_ranges": [f"{a}:{b}" for a, b in parsed],
        "formula_reference_update": "not_guaranteed_by_openpyxl",
    }
    return payload, _render_payload(payload, "deleted_ranges")


def _valid_sheet_name(name: Any) -> str:
    if not isinstance(name, str):
        raise ReadToolError("invalid_sheet_name", "sheet_name must be a string")
    name = name.strip()
    if not name or len(name) > 31 or _SHEET_INVALID_RE.search(name):
        raise ReadToolError("invalid_sheet_name", "invalid Excel worksheet name")
    return name


def _manage_sheet(path: str, arguments: dict[str, Any]) -> tuple[dict[str, Any], str]:
    _require_no_unknown(arguments, {"operation", "sheet_name", "new_name", "index"})
    operation = str(arguments.get("operation") or "").strip().lower()
    if operation not in {"create", "rename", "copy", "move", "hide", "unhide"}:
        raise ReadToolError("invalid_operation", "unsupported sheet operation")
    sheet_name = _valid_sheet_name(arguments.get("sheet_name"))
    with _workbook_lock(path):
        workbook = _open_for_edit(path)
        try:
            if operation == "create":
                if sheet_name in workbook.sheetnames:
                    raise ReadToolError("sheet_exists", "worksheet already exists")
                worksheet = workbook.create_sheet(sheet_name)
            else:
                if sheet_name not in workbook.sheetnames:
                    raise ReadToolError("sheet_not_found", f"worksheet not found: {sheet_name}")
                worksheet = workbook[sheet_name]
                if operation == "rename":
                    new_name = _valid_sheet_name(arguments.get("new_name"))
                    if new_name in workbook.sheetnames and new_name != sheet_name:
                        raise ReadToolError("sheet_exists", "worksheet already exists")
                    worksheet.title = new_name
                elif operation == "copy":
                    new_name = _valid_sheet_name(arguments.get("new_name"))
                    if new_name in workbook.sheetnames:
                        raise ReadToolError("sheet_exists", "worksheet already exists")
                    worksheet = workbook.copy_worksheet(worksheet)
                    worksheet.title = new_name
                elif operation == "move":
                    index = arguments.get("index")
                    if isinstance(index, bool) or not isinstance(index, int):
                        raise ReadToolError("invalid_index", "index must be a zero-based integer")
                    if not 0 <= index < len(workbook._sheets):
                        raise ReadToolError("invalid_index", "index is outside the sheet list")
                    workbook._sheets.remove(worksheet)
                    workbook._sheets.insert(index, worksheet)
                elif operation == "hide":
                    visible = [ws for ws in workbook.worksheets if ws.sheet_state == "visible"]
                    if worksheet.sheet_state == "visible" and len(visible) <= 1:
                        raise ReadToolError(
                            "cannot_hide_only_visible_sheet",
                            "cannot hide the only visible worksheet",
                        )
                    worksheet.sheet_state = "hidden"
                elif operation == "unhide":
                    worksheet.sheet_state = "visible"
            _atomic_save(workbook, path)
            sheets = [{"name": ws.title, "state": ws.sheet_state} for ws in workbook.worksheets]
        finally:
            workbook.close()
    payload = {"status": "success", "operation": operation, "sheets": sheets}
    return payload, _render_payload(payload, "sheets")


def execute_structure_tool(
    name: str, path: str, arguments: dict[str, Any]
) -> tuple[dict[str, Any], str]:
    if not isinstance(arguments, dict):
        raise ReadToolError("invalid_arguments", "arguments must be an object")
    if name == "format_range":
        return _format_range(path, arguments)
    if name == "delete_rows":
        return _delete_axis(path, arguments, axis="rows")
    if name == "delete_columns":
        return _delete_axis(path, arguments, axis="columns")
    if name == "manage_sheet":
        return _manage_sheet(path, arguments)
    raise ReadToolError("unknown_tool", f"unknown structure tool: {name}")
